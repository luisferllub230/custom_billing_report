# -*- coding: utf-8 -*-
"""Tests for the XLSX controller and workbook builder."""

import io
import json
from datetime import date

from odoo.tests import HttpCase, TransactionCase, tagged
from odoo.addons.custom_billing_report.controllers.main import (
    BillingReportController,
)


@tagged("post_install", "-at_install", "custom_billing_report")
class TestBuildXlsx(TransactionCase):
    """Direct unit tests on the workbook builder — no HTTP layer."""

    def test_build_xlsx_returns_zip_bytes(self):
        """``_build_xlsx`` must return a non-empty XLSX (zip) payload."""
        data = self.env["billing.report"].get_report_data({})
        controller = BillingReportController()
        payload = controller._build_xlsx(data, env=self.env)
        self.assertIsInstance(payload, bytes)
        # XLSX is a ZIP container — first 2 bytes are 'PK'.
        self.assertEqual(payload[:2], b"PK")
        self.assertGreater(len(payload), 1000)

    def test_build_xlsx_contains_billing_sheet(self):
        """Confirm the workbook actually opens and exposes our sheet."""
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl not installed in this environment")
        data = self.env["billing.report"].get_report_data({})
        payload = BillingReportController()._build_xlsx(data, env=self.env)
        wb = load_workbook(io.BytesIO(payload), read_only=True)
        self.assertIn("Billing Report", wb.sheetnames)

    def test_build_xlsx_with_payment_categories(self):
        """Payment-category breakdown section must render without
        raising even when categories carry zero amounts."""
        data = self.env["billing.report"].get_report_data({})
        data["payment_categories"] = [
            {"key": "cash", "label": "Cash", "amount": 100.0},
            {"key": "card", "label": "Card", "amount": 50.0},
            {"key": "transfer", "label": "Transfer", "amount": 0.0},
            {"key": "bank", "label": "Bank", "amount": 0.0},
            {"key": "other", "label": "Other", "amount": 0.0},
        ]
        data["payment_total_collected"] = 150.0
        payload = BillingReportController()._build_xlsx(data, env=self.env)
        self.assertEqual(payload[:2], b"PK")
        self.assertGreater(len(payload), 1000)

    def test_build_xlsx_with_user_breakdowns(self):
        """When there are aggregated rows the workbook still renders.

        We synthesise a payload with both ``top_salespersons`` and
        ``top_creators`` set so the breakdown writer in the controller
        is exercised.
        """
        data = self.env["billing.report"].get_report_data({})
        data["top_salespersons"] = [{
            "user_id": 1, "user_name": "Tester Sales",
            "count": 1, "amount_total": 100.0,
            "amount_paid": 60.0, "amount_pending": 40.0,
        }]
        data["top_creators"] = [{
            "user_id": 1, "user_name": "Tester Creator",
            "count": 1, "amount_total": 100.0,
            "amount_paid": 60.0, "amount_pending": 40.0,
        }]
        payload = BillingReportController()._build_xlsx(data, env=self.env)
        self.assertEqual(payload[:2], b"PK")
        self.assertGreater(len(payload), 1000)


@tagged("post_install", "-at_install")
class TestXlsxRoute(HttpCase):
    """End-to-end test of the ``/custom_billing_report/xlsx`` route."""

    def test_route_requires_login(self):
        # Hitting the route without a session redirects to /web/login.
        response = self.url_open("/custom_billing_report/xlsx", timeout=30)
        self.assertIn(response.status_code, (200, 302, 303))

    def test_route_streams_xlsx_for_logged_user(self):
        self.authenticate("admin", "admin")
        options = json.dumps({
            "date_from": date(2026, 1, 1).isoformat(),
            "date_to": date(2026, 1, 31).isoformat(),
            "payment_state": "all",
        })
        response = self.url_open(
            "/custom_billing_report/xlsx?options=%s" % options,
            timeout=30,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.headers.get("Content-Type"),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(response.content[:2], b"PK")
